"""Render measured outcomes into Markdown, standalone HTML and an Excel workbook."""

from __future__ import annotations

import csv
import html
import json
import math
import re
import statistics
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import markdown

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "evals/reports/comprehensive"
DOC = ROOT / "docs/IMPLEMENTATION_AND_MODEL_EVALUATION.md"
DOMAINS = [
    "clarity",
    "customer_need",
    "market_potential",
    "feasibility",
    "competition",
    "differentiation",
]


def read(name, default=None):
    path = OUT / name
    return (
        json.loads(path.read_text(encoding="utf-8-sig"))
        if path.exists()
        else default or {}
    )


def percent(value):
    return "Not measured" if value is None else f"{float(value) * 100:.2f}%"


def cell(value):
    if value is None:
        return "Unassessed"
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    return html.escape(str(value)).replace("|", "&#124;").replace("\n", "<br>")


def expectation(assertion):
    kind = assertion["kind"]
    if kind == "present":
        return "Field present with source evidence"
    if kind == "missing":
        return "Value null; status missing"
    if kind in {"score_none", "headline_null"}:
        return "Null / withheld"
    if kind == "ranking_false":
        return False
    if kind == "identity_absent":
        return "Exclude " + str(assertion.get("expected")) + " from scored data"
    if kind == "low_domain":
        return (
            "<= "
            + str(assertion["expected"])
            + " or cautious abstention (reported separately)"
        )
    if kind == "money":
        return "INR " + str(assertion["expected"])
    return assertion.get("expected")


def table(headers, rows):
    return (
        "\n".join(
            [
                "| " + " | ".join(headers) + " |",
                "| " + " | ".join("---" for _ in headers) + " |",
            ]
            + ["| " + " | ".join(cell(v) for v in row) + " |" for row in rows]
        )
        + "\n"
    )


def p95(values):
    return sorted(values)[max(0, math.ceil(0.95 * len(values)) - 1)] if values else None


def main():
    audit = read("task-audit.json")
    runtime = read("runtime.json")
    business = read("business.json")
    cases = business.get("cases", [])
    expected = json.loads(
        (ROOT / "evals/benchmarks/comprehensive-business-cases.json").read_text(
            encoding="utf-8"
        )
    )["cases"]
    order = {c["id"]: i for i, c in enumerate(expected)}
    cases = sorted(cases, key=lambda c: order[c["case"]])
    components = {
        name: read(name + ".json")
        for name in ["language", "embedding", "translation", "ocr", "speech", "odia"]
    }
    assertions = [{"case": c["case"], **a} for c in cases for a in c["checks"]]
    completed = [c for c in cases if c["status"] == "completed"]
    timings = [c["seconds"] for c in cases]
    measures = [m for c in cases for m in c.get("measurements", [])]
    total_tokens = sum(int(m.get("output_tokens", 0)) for m in measures)
    total_generation = sum(float(m.get("seconds", 0)) for m in measures)
    money = [a for a in assertions if a["kind"] == "money"]
    low = [a for a in assertions if a["kind"] == "low_domain"]
    summary = {
        "evaluated_ist": datetime.now(
            timezone(timedelta(hours=5, minutes=30))
        ).isoformat(),
        "data_scope": "Authored synthetic fixtures; no independent population labels or success outcomes",
        "task_status_counts": audit.get("status_counts"),
        "fresh_api_tests": audit.get("fresh_api_regression"),
        "business_cases_requested": len(expected),
        "business_cases_measured": len(cases),
        "business_completed": len(completed),
        "business_assertions_passed": sum(a["passed"] for a in assertions),
        "business_assertions_total": len(assertions),
        "business_all_checks_passed_cases": sum(c["passed"] for c in cases),
        "money_exact_passed": sum(a["passed"] for a in money),
        "money_exact_total": len(money),
        "low_domain_requested": len(low),
        "low_domain_cautious_abstentions": sum(a.get("abstained", False) for a in low),
        "low_domain_actual_low_ratings": sum(
            a["passed"] and not a.get("abstained", False) for a in low
        ),
        "low_domain_overly_high_ratings": sum(not a["passed"] for a in low),
        "end_to_end_p50_seconds": statistics.median(timings) if timings else None,
        "end_to_end_p95_seconds": p95(timings),
        "qwen_peak_allocated_vram_bytes": max(
            (int(m.get("peak_vram_bytes", 0)) for m in measures), default=0
        ),
        "qwen_generation_tokens_per_second": total_tokens / total_generation
        if total_generation
        else None,
        "component_status": {
            n: d.get("status", "not run") for n, d in components.items()
        },
        "human_quality_validated": False,
        "deployed": False,
    }
    (OUT / "summary.json").write_text(
        json.dumps(summary, indent=2) + "\n", encoding="utf-8"
    )
    byid = {c["case"]: c for c in cases}
    model_rows = []
    model_rows.append(
        [
            "Qwen3-VL-2B-Instruct",
            "Business facts + six domain judgments",
            f"{len(cases)}/{len(expected)} cases measured",
            f"{summary['business_assertions_passed']}/{len(assertions)} scenario checks; {summary['money_exact_passed']}/{len(money)} money facts",
            f"End-to-end p50 {summary['end_to_end_p50_seconds']}; p95 {summary['end_to_end_p95_seconds']} seconds",
            "Actual HTTP/API/Postgres + owned GPU; this is not success-prediction accuracy",
        ]
    )
    for name, label, task in [
        ("language", "py3langid", "Language identification"),
        ("embedding", "MiniLM-L6-v2", "384D semantic retrieval"),
        ("translation", "NLLB-200-distilled-600M", "Regional translation"),
        ("ocr", "PaddleOCR-VL-1.5", "Page OCR"),
        ("speech", "Whisper-large-v3-turbo", "English synthetic speech transcription"),
        ("odia", "Whisper-large-v3-vaani-odia", "Declared Odia speech"),
    ]:
        d = components[name]
        metrics = d.get("metrics", {})
        if name == "language":
            quality = f"Code match {percent(metrics.get('exact_selected_code_accuracy'))}; auto coverage {percent(metrics.get('automatic_detection_coverage'))}"
        elif name == "embedding":
            quality = "; ".join(
                k + " " + percent(v) for k, v in d.get("retrieval_metrics", {}).items()
            )
        elif name == "translation":
            quality = f"Mean sentence chrF++ {metrics.get('mean_sentence_chrfpp')}; number preservation {percent(metrics.get('number_preservation_rate'))}"
        elif name == "ocr":
            quality = f"WER {percent(metrics.get('micro_wer'))}; CER {percent(metrics.get('micro_cer'))}; numbers {percent(metrics.get('number_preservation_rate'))}"
        elif name == "speech":
            quality = f"WER {percent(metrics.get('micro_wer'))}; real-time factor {metrics.get('real_time_factor_warm')}"
        else:
            quality = (
                "WER/CER not measurable: no independent reference; Odia script present "
                + str(d.get("odia_script_present"))
            )
        model_rows.append(
            [
                label,
                task,
                d.get("status", "Not run"),
                quality,
                f"Load {d.get('load_seconds', 'n/a')} s; total {d.get('total_seconds', 'n/a')} s; peak {round(d.get('peak_vram_bytes', 0) / 1024**2, 1)} MiB",
                metrics.get("note")
                or d.get("reason")
                or "Small authored challenge only; not independent population validation",
            ]
        )
    embedding = components["embedding"]
    candidate = embedding.get("experimental_candidate_metrics", {})
    model_rows.append(
        [
            "Experimental logistic relevance head",
            "QA relevance only; not served",
            "Rejected for serving",
            f"Synthetic strict accuracy {percent(candidate.get('strict_accuracy'))}; macro-F1 {candidate.get('macro_f1')}; irrelevant precision {percent(candidate.get('irrelevant_precision'))}",
            "CPU over frozen256-token QA embeddings",
            "No new training. 102 unreviewed original training pairs; new16-pair challenge cannot promote it.",
        ]
    )
    model_rows.append(
        [
            "Optional Groq GPT-OSS explanation",
            "Explains already-saved scores",
            "Not configured on this machine",
            "No new quality/accuracy measurement",
            "Historical20B English/Hindi1.194/1.334 s; not remeasured",
            "Cannot change scores. Existing provider report is historical, not current credentials/quota proof.",
        ]
    )
    sources = [
        "Metric implementation: evals/comprehensive_metrics.py (five known-answer tests pass).",
        "Full original organizer corpus has no completed independent annotation set; T20 remains blocked.",
        "New measurements use RTX3050 6144MiB. Older RTX5050/8151MiB numbers are marked historical.",
        "Local migrations001-008 applied only to loopback test DB. Existing app credentials/processes were preserved.",
    ]
    sections = []

    def add(text):
        sections.append(text.strip() + "\n")

    add("# Implementation and model evaluation — 18 October 2026")
    add(
        f"Generated {summary['evaluated_ist']}. Code audited at `{runtime.get('git_head')}` on branch `{runtime.get('branch')}`. This is an evaluation of implemented behavior, not a public release or a validated funding/business-success model."
    )
    add("## Read this first")
    add(
        table(
            ["Question", "Measured answer"],
            [
                [
                    "How much engineering work is done?",
                    f"All29 tasks audited: {audit.get('status_counts', {}).get('done', 0)} done, {audit.get('status_counts', {}).get('blocked', 0)} blocked. Many blocked tasks have working features but incomplete required validation. This is not a product-completion percentage.",
                ],
                [
                    "Do ordinary code checks pass?",
                    f"Fresh API regression: {audit.get('fresh_api_regression', {}).get('tests')} tests, {audit.get('fresh_api_regression', {}).get('failures')} failures, {audit.get('fresh_api_regression', {}).get('skipped')} skips; suite {audit.get('fresh_api_regression', {}).get('suite_seconds')} seconds. Five new metric tests pass.",
                ],
                [
                    "Were the models actually run?",
                    f"{len(cases)} full HTTP/API/PG business cases plus component results below. Auth operator in this harness is a test fixture. Model outputs are actual inference, not canned assessment fixtures.",
                ],
                [
                    "Can negative businesses reliably receive negative judgments?",
                    "Not reliably demonstrated. Read the weak-demand and funding-gap counterexamples. Correct financial warnings and safe abstention do not prove correct grading.",
                ],
                [
                    "Can a missing answer be treated as a bad business?",
                    "No. Unknown facts/domains remain null; clarification is separate from evidenced weakness.",
                ],
                [
                    "Is the application publicly deployed?",
                    "No. Local capability and historical Supabase checks do not establish a public release.",
                ],
                ["Which current machine?", runtime.get("gpu")],
                [
                    "Current normal configuration?",
                    f"Root app database configured: {runtime.get('normal_app_database_configured_in_root')}; optional Groq configured: {runtime.get('optional_groq_configured')}. Evaluation uses separate loopback test PG, not hosted applicant data.",
                ],
            ],
        )
    )
    add("## What positive, negative and unknown mean")
    add(
        "The app does not output a learned binary business-success label. It extracts supplied facts, proposes six rubric ratings and calculates totals. A well-described proposal may deserve stronger supported dimensions; an explicitly flawed proposal should expose the flaw and lower the relevant dimension or abstain with a reason. Missing information should be unknown, not zero. A fictional applicant saying that demand is strong is still a self-report, not proof of real market success."
    )
    add(
        "The frozen test expectations were authored before inference. They are useful regression targets, not independently reviewed business labels. `low_domain` accepts null as cautious abstention and reports that separately; it does not count null as a confirmed correct negative rating. A 50/100 score here is a proposed rubric calculation, not a50% success probability."
    )
    add("## How the implemented system works")
    add(
        "1. Supabase Auth and database-backed roles/cohorts authorize private submission and review. The UI supports guided text, PDF and recording/upload inputs.\n2. Durable Postgres stages preserve the original and per-stage progress. Digital PDFs use their text layer; scans use PaddleOCR; speech uses Whisper with actual sample time ranges.\n3. Language identification and NLLB prepare English evidence, preserving originals and uncertainty.\n4. Qwen suggests source quotations and business fields. Exact-span rules recover explicit statements; typed schemas reject invented quotes/numbers, embedded instructions and identity-based scoring. Conflicts remain unresolved alternatives.\n5. Qwen makes domain judgments using eligible business fields. The API recomputes Decimal totals, coverage and completeness; it may withhold headline/ranking for uncertain sources, critical conflicts or insufficient evidence.\n6. Separate deterministic financial checks compare funding, sales-based revenue and operating surplus. They currently do not automatically downgrade the model rating.\n7. Frozen MiniLM384D vectors support scoped similar ideas. SQL aggregates drive four comparison views. Human review/corrections, linked clarifications, versioned reports and a separate what-if planner complete the local workflow."
    )
    add("## All29 tasks: implementation, method, performance and gaps")
    add(
        table(
            [
                "Task",
                "Tracker state",
                "Implemented behavior",
                "How / code",
                "Evidence / performance",
                "Remaining limits",
            ],
            [
                [
                    r["id"] + " " + r["task"],
                    r["tracker_status"],
                    r["implemented"],
                    r["how"],
                    r["evidence_performance"],
                    r["limitations"],
                ]
                for r in audit.get("tasks", [])
            ],
        )
    )
    add("## Model inventory and measured quality")
    add(
        table(
            [
                "Model",
                "Role",
                "Current evaluation",
                "Quality measure",
                "Performance",
                "Interpretation",
            ],
            model_rows,
        )
    )
    catalog = read("model-catalog.json")
    add(
        "Legacy AI4Bharat models below are not selected live adapters. They were access-checked only, not falsely counted as inference:"
    )
    add(
        table(
            [
                "Artifact",
                "Repository",
                "Cached weights",
                "Config HTTP status",
                "Selected",
            ],
            [
                [
                    r["artifact"],
                    r["id"],
                    r["weight_files"],
                    r.get("config_http_status", "not probed"),
                    r["selected"],
                ]
                for r in catalog.get("models", [])
                if not r["selected"]
            ],
        )
    )
    add("## Business examples: expected versus actual")
    add(
        f"{len(completed)}/{len(expected)} requested scenarios completed processing; {summary['business_all_checks_passed_cases']}/{len(cases)} measured cases passed all their frozen assertions. Checks passed {summary['business_assertions_passed']}/{len(assertions)}. These are assertion counts, not a claim of overall model/business-success accuracy. Gold money facts matched {summary['money_exact_passed']}/{len(money)} including amount/currency; frequency/unit correctness is also exercised by the financial checks."
    )
    add(
        table(
            [
                "Case",
                "Perspective",
                "Processing",
                "Seconds",
                "Passed checks",
                "Provisional /100",
                "Headline",
                "Coverage",
                "Completeness",
                "Reason withheld",
            ],
            [
                [
                    c["case"],
                    c["perspective"],
                    c["status"],
                    c["seconds"],
                    f"{sum(a['passed'] for a in c['checks'])}/{len(c['checks'])}",
                    c.get("scores", {}).get("provisional_potential"),
                    c.get("scores", {}).get("headline_potential"),
                    percent(c.get("scores", {}).get("evidence_coverage")),
                    str(c.get("scores", {}).get("completeness_percent", "?")) + "%",
                    c.get("scores", {}).get("withheld_reason"),
                ]
                for c in cases
            ],
        )
    )
    add("### Six domain judgments")
    add(
        "Scale:0 evidenced weakness,1 weak support,2 plausible/limited,3 supported self-report;4 requires independent verification and is unavailable to this self-reported worker. Unassessed means null, not zero."
    )
    add(
        table(
            ["Case"] + DOMAINS,
            [
                [c["case"]]
                + [
                    next(
                        (
                            d["rating"]
                            for d in c.get("domains", [])
                            if d["domain"] == name
                        ),
                        None,
                    )
                    for name in DOMAINS
                ]
                for c in cases
            ],
        )
    )
    add("### The important counterexamples")
    add(
        f"Among {len(low)} predeclared negative-domain expectations, {summary['low_domain_actual_low_ratings']} actually received a low rating, {summary['low_domain_cautious_abstentions']} abstained, and {summary['low_domain_overly_high_ratings']} received a rating above the expected ceiling. These are correlated authored checks, not independent expert-labelled businesses."
    )
    failed = [a for a in assertions if not a["passed"]]
    add(
        table(
            ["Case", "Check", "Expected", "Observed", "Meaning"],
            [
                [
                    a["case"],
                    a["kind"] + (":" + a["field"] if a.get("field") else ""),
                    expectation(a),
                    a.get("observed"),
                    "Does not meet the predeclared expectation; retained without tuning the application.",
                ]
                for a in failed
            ],
        )
    )
    if "N01" in byid:
        add(
            "**Loss example N01:** supplied monthly revenue₹10000 and expenses₹18000 imply an operating surplus of **−₹8000**. The financial check detects this. A null feasibility dimension is cautious abstention; it is not a negative model rating. Compare its provisional total with the strong examples before interpreting the score."
        )
    if "N03" in byid:
        rating = next(
            (d for d in byid["N03"].get("domains", []) if d["domain"] == "feasibility"),
            {},
        )
        add(
            f"**Funding example N03:** ₹50000 needed, ₹1000 own funds, ₹0 loan. The rule flags a discrepancy, but the model feasibility rating is **{rating.get('rating')} /4**. Its actual reason: {cell(rating.get('reason'))}. A critical-risk warning has not been converted into a reliably low rubric judgment."
        )
    if "P02" in byid:
        profit = byid["P02"].get("business", {}).get("profit", {}).get("value")
        add(
            f"**Period-loss example P02:** the source explicitly states monthly profit. Extracted profit is `{cell(profit)}`. A missing period can prevent checking an otherwise supplied profit even when the amount is correct; see the failed financial-consistency assertion."
        )
    add("### Identity and language/grammar perspectives")
    if "B01" in byid and "B02" in byid:
        left, right = byid["B01"], byid["B02"]
        domains_left = {d["domain"]: d["rating"] for d in left.get("domains", [])}
        domains_right = {d["domain"]: d["rating"] for d in right.get("domains", [])}
        differences = {
            k: [domains_left.get(k), domains_right.get(k)]
            for k in DOMAINS
            if domains_left.get(k) != domains_right.get(k)
        }
        add(
            f"B01/B02 contain the same business and vary only irrelevant name/gender/education prefaces. Numeric domain differences: `{cell(differences)}`. Names should be absent from scored facts. This one authored pair cannot establish fairness; extracted quotations, segment IDs and model variation may also affect results."
        )
    add(
        "G01 checks imperfect English with concrete business facts. H01 checks Hindi processing. A missed fact is recorded, not explained away as a bad business. Original/English source uncertainty may withhold a headline even for an otherwise supported proposal."
    )
    add("## Actual performance: what these timings include")
    add(
        table(
            ["Measure", "Result", "Scope"],
            [
                [
                    "Full business case p50 seconds",
                    summary["end_to_end_p50_seconds"],
                    "Submit -> stages -> persisted assessment, localhost HTTP/test PG; includes queue/poll/model process loads",
                ],
                [
                    "Full business case p95 seconds",
                    summary["end_to_end_p95_seconds"],
                    "Nearest-rank p95 from this small mixed-complexity scenario set, not an SLA",
                ],
                [
                    "Qwen allocated peak MiB",
                    round(summary["qwen_peak_allocated_vram_bytes"] / 1024**2, 2),
                    "Torch peak allocation, not all reserved/driver/system GPU memory",
                ],
                [
                    "Qwen recorded published-stage tokens/second",
                    round(summary["qwen_generation_tokens_per_second"], 2)
                    if summary["qwen_generation_tokens_per_second"]
                    else None,
                    "Recorded published-stage output tokens divided by their generation seconds; excludes failed-stage measurements, weight load and other stages",
                ],
                [
                    "Historical API metadata p50/p95",
                    "86.335 /212.115 ms",
                    "4988 requests/303s,50 paced local sessions; no LLM inference",
                ],
                [
                    "Historical heavy chart p95",
                    "812.092 ms",
                    "Part of the older local metadata experiment, not current-hardware model latency",
                ],
                [
                    "Historical one-second diagnostic p95",
                    "1944.633 ms",
                    "Faster request rate missed latency target; zero errors",
                ],
                [
                    "Historical real Supabase saved-read latency",
                    "19.683 ->4.519 s",
                    "One account, one sample, three requests ->one snapshot; no inference; not a distribution",
                ],
            ],
        )
    )
    add(
        "Fresh model runs use the current6GB GPU. Earlier local evidence used an8GB RTX5050. Cached model process load is not an uncached cold-service benchmark. Public model provisioning ran in the background during some business tests, so these measurements are useful observed latencies, not fully isolated scientific performance trials. No public free service was load-tested."
    )
    add("## Component results and references")
    lang = components["language"]
    add("### Language identification")
    add(
        "Every test used preferred language `und`, so declared hints cannot masquerade as detection accuracy. Internal detector confidence is not a calibrated business/source-quality probability."
    )
    add(
        table(
            [
                "Case",
                "Expected",
                "Selected",
                "Provenance",
                "Detector confidence",
                "Confirmation needed",
                "Match",
            ],
            [
                [
                    r["id"],
                    r["expected_language"],
                    r["result"]["code"],
                    r["result"]["provenance"],
                    r["result"].get("model_confidence"),
                    r["result"]["needs_confirmation"],
                    r["correct"],
                ]
                for r in lang.get("cases", [])
            ],
        )
    )
    trans = components["translation"]
    add("### Translation: meaning, figures and negation")
    add(
        "**Inspection of T02:** numbers were preserved (10000/18000), but the output omitted the explicit Hindi statement that the business is making a loss. The literal required-term check also fails because the output formats those numbers with commas. Thus the90% term-check result partly reflects formatting, while the omitted loss clause is a separate meaningful defect. This is assistant inspection of an authored fixture, not independent human validation. Number preservation alone is insufficient."
    )
    add(
        "The references below were authored for this challenge, not supplied by an independent fluent reviewer. chrF++ measures character/word n-gram agreement; a valid paraphrase can score lower and a lexical match does not prove factual correctness. [SacreBLEU primary documentation](https://github.com/mjpost/sacrebleu)."
    )
    add(
        table(
            [
                "Case",
                "Direction",
                "Source",
                "Actual output",
                "Reference",
                "chrF++ /100",
                "Numbers preserved",
                "Required terms",
            ],
            [
                [
                    r["id"],
                    r["source_language"] + " -> " + r["target_language"],
                    r["source"],
                    r.get("output", r.get("error_type")),
                    r["reference"],
                    round(r["chrfpp"], 2) if "chrfpp" in r else None,
                    r.get("numbers_preserved"),
                    r.get("required_terms_present"),
                ]
                for r in trans.get("cases", [])
            ],
        )
    )
    add("### OCR: known synthetic references")
    add(
        table(
            [
                "Case",
                "Image type",
                "Reference",
                "Actual OCR",
                "WER",
                "CER",
                "Figures preserved",
                "Seconds",
            ],
            [
                [
                    r["id"],
                    r["style"],
                    r["reference"],
                    r.get("output", r.get("error_type")),
                    percent(r.get("wer")),
                    percent(r.get("cer")),
                    r.get("numbers_preserved"),
                    r["seconds"],
                ]
                for r in components["ocr"].get("cases", [])
            ],
        )
    )
    add(
        "Rendered script font is typed text, not authentic handwriting. Original handwritten/overwritten corpus quality remains unvalidated; earlier full-scan processing included unreadable/capped pages. Normalized CER keeps punctuation and collapses whitespace; WER removes punctuation/symbols, case-folds and tokenizes whitespace. Numeric preservation is checked separately."
    )
    speech = components["speech"]
    add("### Speech: reference transcript versus recognized text")
    add(
        table(
            [
                "Case",
                "Audio seconds",
                "Inference seconds",
                "WER",
                "Reference words",
                "Edits",
                "Figures preserved",
                "Actual transcript",
            ],
            [
                [
                    r["id"],
                    round(r["audio_seconds"], 2),
                    r["seconds"],
                    percent(r["wer"]),
                    r["reference_words"],
                    r["word_edits"],
                    r["numbers_preserved"],
                    r["output"],
                ]
                for r in speech.get("cases", [])
            ],
        )
    )
    add(
        "These recordings use one local Windows synthetic English voice. WER penalizes textual differences, including some spoken-number/spelling conventions. It cannot establish real regional, accented, noisy or diverse-speaker accuracy. References are in input/audio/*.transcript.txt. WER=(substitutions+deletions+insertions)/reference words and can exceed100%; lower is better. It must not be presented as a success probability. [Primary WER definition](https://huggingface.co/spaces/evaluate-metric/wer/blob/main/README.md)."
    )
    add(
        "Silence behavior: "
        + cell(speech.get("silence", {}))
        + ". Compare raw ASR hallucination with the production RMS-based silent-chunk rejection; rejection is a pipeline safeguard, not accurate words from silence."
    )
    add(
        "Odia functional smoke: "
        + cell(
            {
                k: v
                for k, v in components["odia"].items()
                if k
                in [
                    "status",
                    "audio_seconds",
                    "seconds",
                    "load_seconds",
                    "transcript_characters",
                    "odia_script_present",
                    "wer",
                    "cer",
                    "reason",
                ]
            }
        )
        + ". Private original/transcript are ignored, not republished into this report."
    )
    add("### MiniLM retrieval and relevance experiment")
    add(
        table(
            [
                "Query",
                "Relevant IDs",
                "Ranking (top5)",
                "Recall@1",
                "Recall@3",
                "Recall@5",
                "Reciprocal rank",
                "nDCG@5",
            ],
            [
                [
                    r["id"] + " " + r["text"],
                    r["relevant_ids"],
                    r["ranked_ids"][:5],
                    r["recall_at_1"],
                    r["recall_at_3"],
                    r["recall_at_5"],
                    r["reciprocal_rank"],
                    r["ndcg_at_5"],
                ]
                for r in embedding.get("retrieval_cases", [])
            ],
        )
    )
    add(
        "This is four authored queries over ten known synthetic ideas, not production annotated-retrieval accuracy. Recall@1 can be50% when two ideas are relevant even if the top result is correct. Production RLS/current-revision/family filtering is separately covered by API tests."
    )
    add(
        table(
            [
                "System",
                "Strict accuracy",
                "Coverage",
                "Selective accuracy",
                "Macro-F1",
                "Irrelevant precision",
                "Irrelevant recall",
            ],
            [
                [
                    name,
                    percent(metrics.get("strict_accuracy")),
                    percent(metrics.get("coverage")),
                    percent(metrics.get("selective_accuracy")),
                    metrics.get("macro_f1"),
                    percent(metrics.get("irrelevant_precision")),
                    percent(metrics.get("irrelevant_recall")),
                ]
                for name, metrics in [
                    (
                        "Serving rules/cosine baseline",
                        embedding.get("baseline_metrics", {}),
                    ),
                    ("Rejected logistic candidate", candidate),
                ]
            ],
        )
    )
    add(
        "Strict accuracy counts abstentions as not correct. Selective accuracy uses only decided cases and must be read with coverage. Precision is not measured when no irrelevant answers are predicted. Candidate probabilities are uncalibrated. The current baseline does not label low cosine as irrelevant; it abstains. The candidate remains rejected regardless of this small synthetic challenge."
    )
    add(
        table(
            ["Case", "Question", "Answer", "Gold", "Baseline", "Candidate"],
            [
                [
                    r["id"],
                    r["question"],
                    r["answer"],
                    ["relevant", "irrelevant"][r["gold"]],
                    {0: "relevant", 1: "irrelevant", -1: "uncertain"}[r["baseline"]],
                    {0: "relevant", 1: "irrelevant", -1: "uncertain"}[r["candidate"]],
                ]
                for r in embedding.get("qa_cases", [])
            ],
        )
    )
    add("## What should be improved next")
    add(
        table(
            ["Priority", "Change to investigate", "Evidence / why"],
            [
                [
                    "1",
                    "Calibrate negative-demand and unfunded-feasibility judgments using reviewed examples; explicitly connect traceable arithmetic/risk to the proposed rubric.",
                    "Negative survey/funding cases can receive moderate dimensions. Do not auto-change ratings during an evaluation to obtain a green report.",
                ],
                [
                    "2",
                    "Preserve monetary periods/units and improve short/rough-English extraction.",
                    "Correct amounts can retain unknown periods, preventing a promised consistency check; G01 lost investment, H01 lost customers.",
                ],
                [
                    "3",
                    "Audit source-review gates and explain withheld headline separately from potential.",
                    "Strong and weak examples may both show a withheld headline; that alone is not semantic discrimination.",
                ],
                [
                    "4",
                    "Create independent English/Hindi annotations with real business families and regional reference transcripts.",
                    "Authored fixtures and one identity pair do not support population accuracy or fairness.",
                ],
                [
                    "5",
                    "Validate genuine handwriting/noisy speech and real retrieval judgments.",
                    "Typed/script-font OCR, clean synthetic speech and a tiny gallery are limited proxies.",
                ],
                [
                    "6",
                    "Retain rejected relevance candidate and current advisory status until proper class/group coverage is reviewed.",
                    "This experiment does not train a success/funding predictor or validate a foundation model.",
                ],
            ],
        )
    )
    add("## Full per-case evidence")
    for case in cases:
        add("### " + case["case"] + " — " + case["perspective"])
        add("Expected: " + cell(case["expected_interpretation"]))
        add("Authored input:\n\n```text\n" + case["source"] + "\n```")
        add(
            table(
                ["Assertion", "Expected", "Observed", "Passed", "Cautious abstention"],
                [
                    [
                        a["kind"] + (":" + a["field"] if a.get("field") else ""),
                        expectation(a),
                        a.get("observed"),
                        a["passed"],
                        a.get("abstained", False),
                    ]
                    for a in case["checks"]
                ],
            )
        )
        add(
            table(
                ["Domain", "Rating", "Actual model reason"],
                [
                    [d["domain"], d["rating"], d["reason"]]
                    for d in case.get("domains", [])
                ],
            )
        )
        add(
            "Full synthetic assessment/economics JSON: `evals/reports/comprehensive/cases/"
            + case["case"]
            + ".json`."
        )
    recheck = read("recheck-F01/business.json")
    if recheck.get("cases"):
        row = recheck["cases"][0]
        add("### Diagnostic repeat of F01")
        add(
            "The unchanged F01 input was repeated once solely to inspect the failure stage. It is excluded from the original14-case/99-check aggregate. Repeat status: "
            + cell(row["status"])
            + ", seconds "
            + str(row["seconds"])
            + "."
        )
        add(
            table(
                ["Stage", "Status", "Error", "Attempts"],
                [
                    [s["stage_key"], s["status"], s["error_code"], s["attempts"]]
                    for s in row.get("stage_states", [])
                ],
            )
        )
    add("## Reproduction and evidence boundaries")
    add(
        "Normal-worker readiness is separate: the preserved .venv-model here still has Transformers4.57.6/hub0.36.0. These new measurements use the isolated evaluation recipe Transformers5.0.0/hub1.33.0, recorded in runtime.json. In particular, a successful Paddle run in this environment does not prove the older normal worker environment is ready. Current root application DB/Groq settings were unavailable; no hosted deployment or normal-app setup repair is claimed."
    )
    add("\n".join("- " + s for s in sources))
    add(
        "The isolated .venv-eval uses the existing read-only Torch/model package directory via a .pth file, with the pinned evaluation packages installed locally. Existing .venv/.venv-model, ignored root/frontend credentials and the running app were not replaced. Public model weights were explicitly downloaded/hashed; inference uses pinned local snapshots. Legacy gated model config requests returned401; optional Groq is not configured here. No applicant data was submitted to an external model provider."
    )
    add(
        "```powershell\n# Use the existing loopback test DSN only; never run fixtures against Supabase.\n.venv-eval/Scripts/python.exe -B -m pytest -c apps/api/pyproject.toml tests/api -q\n.venv-eval/Scripts/python.exe -B -m pytest tests/evals/test_comprehensive_metrics.py -q\n.venv-eval/Scripts/python.exe -B scripts/evaluate_business_scenarios.py\n# Run these sequentially AFTER the business harness; one GPU model at a time.\n.venv-eval/Scripts/python.exe -B scripts/evaluate_model_components.py language\n.venv-eval/Scripts/python.exe -B scripts/evaluate_model_components.py embedding\n.venv-eval/Scripts/python.exe -B scripts/evaluate_model_components.py translation\n.venv-eval/Scripts/python.exe -B scripts/evaluate_model_components.py ocr\n.venv-eval/Scripts/python.exe -B scripts/evaluate_model_components.py speech\n.venv-eval/Scripts/python.exe -B scripts/evaluate_model_components.py odia\n.venv-eval/Scripts/python.exe -B evals/task_implementation_audit.py\n.venv-eval/Scripts/python.exe -B evals/render_comprehensive_report.py\n```"
    )
    add(
        "Frozen cases and expectations are under evals/benchmarks/. Runtime/package/model revisions and dataset digests are in runtime.json, model-catalog.json and each result JSON. Previous failures remain in their artifacts. Reports and workbook are local; nothing was pushed or published. NLLB is CC-BY-NC-4.0 and its upstream card limits intended use to research; a commercial deployment needs separate model/license review. [Upstream NLLB model card](https://huggingface.co/facebook/nllb-200-distilled-600M)."
    )
    document = "\n".join(sections)
    DOC.write_text(document, encoding="utf-8")
    rendered = markdown.markdown(document, extensions=["tables", "fenced_code", "toc"])
    page = (
        """<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Implementation and model evaluation</title><style>body{font:15px/1.6 system-ui,sans-serif;color:#20352c;background:#f7f8f4;margin:0}main{max-width:1400px;margin:auto;padding:24px}h1,h2,h3{line-height:1.3}h2{margin-top:44px;border-top:1px solid #bccbbd;padding-top:22px}table{border-collapse:collapse;display:block;overflow:auto;max-width:100%;background:white;margin:18px 0}th,td{padding:10px;border:1px solid #d9e1d5;vertical-align:top;min-width:100px;max-width:500px;overflow-wrap:anywhere}th{background:#e3eddd;text-align:left;position:sticky;top:0}tr:nth-child(even){background:#f3f6ee}pre{white-space:pre-wrap;background:#edf1e7;padding:16px}code{font-size:13px;overflow-wrap:anywhere}a{color:#216445}header{position:sticky;top:0;z-index:10;background:#1f4936;color:white;padding:14px 24px;box-shadow:0 2px 8px #0002}input{font:inherit;padding:8px;min-width:280px;max-width:85%;border:0;border-radius:4px}header small{margin-left:15px}button{font:inherit;padding:8px;margin-left:10px;cursor:pointer}@media print{header{display:none}main{max-width:none;padding:0}table{display:table;font-size:9px}th,td{min-width:0;padding:4px}h2{break-before:page}}@media(max-width:600px){main{padding:15px}header{padding:10px}header small{display:block;margin:4px 0}}</style><header><label>Find in result tables <input id="search" type="search" placeholder="e.g. N03, feasibility, loss"></label><button onclick="window.print()">Print / PDF</button><small>Local report. Unassessed is not zero; assertions are not population accuracy.</small></header><main>"""
        + rendered
        + """</main><script>document.getElementById('search').addEventListener('input',e=>{const q=e.target.value.toLowerCase();document.querySelectorAll('table tbody tr').forEach(row=>{row.hidden=q&&!row.textContent.toLowerCase().includes(q)})})</script></html>"""
    )
    (OUT / "IMPLEMENTATION_AND_MODEL_EVALUATION.html").write_text(
        page, encoding="utf-8"
    )
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, headers, rows):
        ws = workbook.create_sheet(name)
        ws.append(headers)
        for row in rows:
            safe = []
            for value in row:
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                if value is None:
                    value = "Unassessed / not measured"
                if (
                    isinstance(value, str)
                    and value.startswith(("=", "+", "-", "@"))
                    and not re.fullmatch(r"-?\d+(\.\d+)?", value)
                ):
                    value = "'" + value
                safe.append(value)
            ws.append(safe)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="244C37")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[row[0].row].height = 60
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = (
                18 if i == 1 else min(65, 28 + (10 if len(headers) < 6 else 0))
            )
            if any(
                word in headers[i - 1].lower()
                for word in (
                    "source",
                    "reference",
                    "transcript",
                    "reason",
                    "explanation",
                    "implemented",
                    "limits",
                    "expected interpretation",
                    "authored",
                    "answer",
                )
            ):
                ws.column_dimensions[get_column_letter(i)].width = 65
        if name == "START HERE":
            ws.column_dimensions["A"].width = 40
            ws.column_dimensions["B"].width = 100
        for row in ws.iter_rows(min_row=2):
            lines = max(
                (
                    math.ceil(
                        len(str(c.value or ""))
                        / max(8, ws.column_dimensions[c.column_letter].width - 3)
                    )
                    + str(c.value or "").count("\n")
                    for c in row
                ),
                default=1,
            )
            ws.row_dimensions[row[0].row].height = min(260, max(32, 15 * lines + 8))
        for i, header in enumerate(headers, 1):
            if header in {
                "WER",
                "CER",
                "Recall@1",
                "Recall@3",
                "Recall@5",
                "MRR",
                "nDCG@5",
                "Coverage",
            }:
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    c = row[0]
                    if isinstance(c.value, (float, int)) and not isinstance(
                        c.value, bool
                    ):
                        c.number_format = "0.00%"
        if name == "Assertions":
            for row in ws.iter_rows(min_row=2):
                row[5].fill = PatternFill(
                    "solid", fgColor="DCF0D8" if row[5].value else "F9D9D5"
                )
        if name == "Business overview":
            for row in ws.iter_rows(min_row=2):
                color = "DCF0D8" if row[5].value == row[6].value else "F9D9D5"
                row[5].fill = PatternFill("solid", fgColor=color)
        ws.sheet_view.zoomScale = 85
        return ws

    sheet(
        "START HERE",
        ["Item", "Answer"],
        [[k, v] for k, v in summary.items()]
        + [
            [
                "Meaning",
                "Scenario assertion scores are not population accuracy. Null is unassessed, not a bad business. Read negative counterexamples and domain reasons.",
            ],
            ["Detailed readable report", "docs/IMPLEMENTATION_AND_MODEL_EVALUATION.md"],
            [
                "Standalone searchable report",
                "IMPLEMENTATION_AND_MODEL_EVALUATION.html",
            ],
        ],
    )
    sheet(
        "All29 tasks",
        [
            "ID",
            "Task",
            "State",
            "Implemented",
            "How / code",
            "Evidence / performance",
            "Limits",
        ],
        [
            [
                r[k]
                for k in [
                    "id",
                    "task",
                    "tracker_status",
                    "implemented",
                    "how",
                    "evidence_performance",
                    "limitations",
                ]
            ]
            for r in audit.get("tasks", [])
        ],
    )
    sheet(
        "Model summary",
        ["Model", "Role", "Status", "Quality", "Timing / memory", "Limits"],
        model_rows,
    )
    sheet(
        "Business overview",
        [
            "Case",
            "Perspective",
            "Expected",
            "Processing",
            "Seconds",
            "Passed checks",
            "Total checks",
            "Provisional",
            "Headline",
            "Coverage",
            "Completeness",
            "Withheld",
        ],
        [
            [
                c["case"],
                c["perspective"],
                c["expected_interpretation"],
                c["status"],
                c["seconds"],
                sum(a["passed"] for a in c["checks"]),
                len(c["checks"]),
                c.get("scores", {}).get("provisional_potential"),
                c.get("scores", {}).get("headline_potential"),
                c.get("scores", {}).get("evidence_coverage"),
                c.get("scores", {}).get("completeness_percent"),
                c.get("scores", {}).get("withheld_reason"),
            ]
            for c in cases
        ],
    )
    sheet(
        "Assertions",
        ["Case", "Check", "Field", "Expected", "Observed", "Pass", "Abstained"],
        [
            [
                a["case"],
                a["kind"],
                a.get("field"),
                expectation(a),
                a.get("observed"),
                a["passed"],
                a.get("abstained", False),
            ]
            for a in assertions
        ],
    )
    sheet(
        "Inputs",
        ["Case", "Perspective", "Authored source", "Expected interpretation"],
        [
            [c["case"], c["perspective"], c["source"], c["expected_interpretation"]]
            for c in cases
        ],
    )
    sheet(
        "Domain reasons",
        ["Case", "Domain", "Rating", "Reason", "Evidence refs"],
        [
            [
                c["case"],
                d["domain"],
                d["rating"],
                d["reason"],
                d.get("source_segment_ids"),
            ]
            for c in cases
            for d in c.get("domains", [])
        ],
    )
    sheet(
        "Extracted facts",
        ["Case", "Field", "Status", "Value", "Alternatives"],
        [
            [c["case"], f, d["status"], d.get("value"), d.get("alternatives")]
            for c in cases
            for f, d in c.get("business", {}).items()
            if isinstance(d, dict) and "status" in d
        ],
    )
    sheet(
        "Financial checks",
        ["Case", "Kind", "Status", "Formula", "Calculated", "Stated", "Explanation"],
        [
            [
                c["case"],
                d["kind"],
                d["status"],
                d.get("formula"),
                d.get("calculated_amount"),
                d.get("stated_amount"),
                d["explanation"],
            ]
            for c in cases
            for d in c.get("economics", {}).get("checks", [])
        ],
    )
    sheet(
        "Translation",
        [
            "ID",
            "Source language",
            "Target language",
            "Source",
            "Reference",
            "Actual output",
            "chrF++",
            "Numbers",
            "Terms",
            "Seconds",
        ],
        [
            [
                r.get(k)
                for k in [
                    "id",
                    "source_language",
                    "target_language",
                    "source",
                    "reference",
                    "output",
                    "chrfpp",
                    "numbers_preserved",
                    "required_terms_present",
                    "seconds",
                ]
            ]
            for r in trans.get("cases", [])
        ],
    )
    sheet(
        "OCR",
        ["ID", "Style", "Reference", "Actual OCR", "WER", "CER", "Numbers", "Seconds"],
        [
            [
                r.get(k)
                for k in [
                    "id",
                    "style",
                    "reference",
                    "output",
                    "wer",
                    "cer",
                    "numbers_preserved",
                    "seconds",
                ]
            ]
            for r in components["ocr"].get("cases", [])
        ],
    )
    sheet(
        "Speech",
        [
            "ID",
            "Reference",
            "Actual transcript",
            "WER",
            "Words",
            "Edits",
            "Audio seconds",
            "Inference seconds",
        ],
        [
            [
                r.get(k)
                for k in [
                    "id",
                    "reference",
                    "output",
                    "wer",
                    "reference_words",
                    "word_edits",
                    "audio_seconds",
                    "seconds",
                ]
            ]
            for r in speech.get("cases", [])
        ],
    )
    sheet(
        "Language ID",
        ["ID", "Expected", "Source", "Result", "Correct", "Seconds"],
        [
            [
                r["id"],
                r["expected_language"],
                r["text"],
                r["result"],
                r["correct"],
                r["seconds"],
            ]
            for r in lang.get("cases", [])
        ],
    )
    sheet(
        "Retrieval",
        [
            "ID",
            "Query",
            "Relevant",
            "Ranking",
            "Recall@1",
            "Recall@3",
            "Recall@5",
            "MRR",
            "nDCG@5",
            "Seconds",
        ],
        [
            [
                r.get(k)
                for k in [
                    "id",
                    "text",
                    "relevant_ids",
                    "ranked_ids",
                    "recall_at_1",
                    "recall_at_3",
                    "recall_at_5",
                    "reciprocal_rank",
                    "ndcg_at_5",
                    "seconds",
                ]
            ]
            for r in embedding.get("retrieval_cases", [])
        ],
    )
    sheet(
        "Relevance QA",
        [
            "ID",
            "Question",
            "Answer",
            "Gold0relevant1irrelevant",
            "Baseline-1abstain",
            "Rejected candidate",
            "Cosine",
            "Uncalibrated candidate probability",
        ],
        [
            [
                r.get(k)
                for k in [
                    "id",
                    "question",
                    "answer",
                    "gold",
                    "baseline",
                    "candidate",
                    "cosine",
                    "candidate_uncalibrated_irrelevant_probability",
                ]
            ]
            for r in embedding.get("qa_cases", [])
        ],
    )
    sheet("Runtime", ["Key", "Value"], [[k, v] for k, v in runtime.items()])
    sheet(
        "Metric definitions",
        ["Metric", "Meaning"],
        [
            [
                "WER / CER",
                "Edit errors divided by reference words/characters; lower better; may exceed100%; empty reference is not measured.",
            ],
            [
                "chrF++",
                "Lexical character/word n-gram F score0-100; higher lexical agreement, not expert semantic correctness.",
            ],
            [
                "Recall@k",
                "Fraction of authored relevant IDs found among firstk; tiny synthetic gallery only.",
            ],
            [
                "Macro-F1",
                "Mean per-class F1; balanced QA relevance check, not business prediction.",
            ],
            [
                "Strict vs selective accuracy",
                "Strict uses all examples including abstentions as incorrect; selective uses only decisions, so read coverage.",
            ],
            [
                "Scenario pass rate",
                "Passed frozen authored checks /requested checks; no independently reviewed population accuracy.",
            ],
            [
                "Unassessed / null",
                "No sufficiently grounded judgment. Never substitute zero, bad business or failure probability.",
            ],
            [
                "Negative economics",
                "Calculated deficit or discrepancy, separate from model domain ratings. No automatic score mutation.",
            ],
        ],
    )
    target = OUT / "EVALUATION_WORKBOOK.xlsx"
    workbook.save(target)
    checked = load_workbook(target, read_only=True, data_only=False)
    assert (
        checked["All29 tasks"].max_row == 30
        and checked["Assertions"].max_row == len(assertions) + 1
    )
    assert not any(
        c.data_type == "f" for ws in checked for row in ws.iter_rows() for c in row
    )
    checked.close()
    for ws in workbook:
        with (OUT / (re.sub(r"[^a-z0-9]+", "-", ws.title.lower()) + ".csv")).open(
            "w", encoding="utf-8-sig", newline=""
        ) as file:
            writer = csv.writer(file)
            writer.writerows(ws.values)
    print(
        json.dumps(
            {
                "markdown": str(DOC.relative_to(ROOT)),
                "workbook_sheets": len(workbook.sheetnames),
                "business_cases": len(cases),
                "assertions": len(assertions),
                "summary": summary["component_status"],
            }
        )
    )


if __name__ == "__main__":
    main()
